"""Act 3: the Q4 payout spreadsheet two ways, on the SAME export.
  Typical  what most teams start from: SUMIFS on owner, Closed Won, in-quarter, split by split_pct, tiered rate, cap
  Fixed    the same sheet plus one helper column per plan rule (FX, duplicates, renewals 50%, deal desk, split check,
           SPIFF) and a reps tab (proration, employed on payment date)
Writes data/Corvane_Q4_Payouts_Spreadsheet.xlsx (live formulas, Arial; blue = inputs)."""
import csv, os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__)); DATA = os.path.join(HERE, "data")
OUT = os.path.join(DATA, "Corvane_Q4_Payouts_Spreadsheet.xlsx")
F = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True)
BL = Font(name="Arial", size=10, color="0000FF")
W = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD, RULE = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="B45309")
REPS = ["Priya Shah", "Tyler Brooks", "Sarah Chen", "Marcus Reed", "Dan Ortiz", "Ana Lima"]


def rows(n):
    with open(os.path.join(DATA, n), newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def d(s):
    return date(*map(int, s.split("-"))) if s else None


def header(ws, r, labels, fill=HEAD, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(r, start + i, t); c.font = W; c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center")


wb = Workbook()
# ---------------- Inputs
inp = wb.active; inp.title = "Inputs"
INPUTS = [("Quarter start", date(2026, 10, 1), "Plan, Overview"), ("Quarter end", date(2026, 12, 31), "Plan, Overview"),
          ("Quota", 300000, "Plan, Overview"), ("Target incentive", 20000, "Plan, Overview"),
          ("Payment date", date(2027, 1, 15), "Plan, General Terms"), ("Renewal credit", 0.5, "Tom's email, Nov 3"),
          ("SPIFF per new logo", 500, "Slack #sales, Nov 2"), ("SPIFF minimum", 50000, "Slack #sales, Nov 2"),
          ("Max discount without VP", 20, "Plan, Deal Desk")]
for r, (k, v, note) in enumerate(INPUTS, 1):
    inp.cell(r, 1, k).font = B; c = inp.cell(r, 2, v); c.font = BL; inp.cell(r, 3, note).font = F
    c.number_format = "yyyy-mm-dd" if isinstance(v, date) else ("0%" if v == 0.5 else "#,##0")
inp.column_dimensions["A"].width = 26; inp.column_dimensions["B"].width = 14; inp.column_dimensions["C"].width = 24
QS, QE, QT, TI, PAY, REN, SP, SPMIN, DISC = (f"Inputs!$B${i}" for i in range(1, 10))

# ---------------- Export (the CRM export pasted as-is) + helper columns
ex = wb.create_sheet("Export"); deals = rows("deals.csv"); cols = list(deals[0]); N = len(deals) + 1
header(ex, 1, cols)
for r, row in enumerate(deals, 2):
    for c, k in enumerate(cols, 1):
        v = row[k]
        if k == "close_date":
            v = d(v)
        elif k in ("amount", "discount_pct", "split_pct", "owner_pct") and v != "":
            v = float(v)
        cell = ex.cell(r, c, v if v != "" else None); cell.font = F
        if k == "close_date":
            cell.number_format = "yyyy-mm-dd"
# A id, C owner, D stage, E type, F close, G amount, H currency, I discount, J approval, K new_logo, L split_rep, M split_pct, N owner_pct
T_COLS = [("Typical: owner credit", lambda r: f'=IF(AND(D{r}="Closed Won",F{r}>={QS},F{r}<={QE}),G{r}*(1-N(M{r})/100),0)'),
          ("Typical: split credit", lambda r: f'=IF(AND(D{r}="Closed Won",F{r}>={QS},F{r}<={QE}),G{r}*N(M{r})/100,0)')]
# helper columns Q..AB
F_COLS = [("FX week", lambda r: f'=IF(H{r}="USD","",_xlfn.MAXIFS(FX!$A$2:$A$100,FX!$B$2:$B$100,H{r},FX!$A$2:$A$100,"<="&F{r}))'),  # Q
          ("USD rate", lambda r: f'=IF(H{r}="USD",1,SUMIFS(FX!$C$2:$C$100,FX!$A$2:$A$100,Q{r},FX!$B$2:$B$100,H{r}))'),             # R
          ("First copy?", lambda r: f'=COUNTIF($A$2:A{r},A{r})=1'),                                                                  # S
          ("Renewal factor", lambda r: f'=IF(E{r}="Renewal",{REN},1)'),                                                             # T
          ("Deal desk OK?", lambda r: f'=OR(N(I{r})<={DISC},J{r}<>"")'),                                                            # U
          ("Owner %", lambda r: f'=IF(L{r}="",100,IF(N{r}="",100-M{r},N{r}))'),                                                     # V
          ("Split adds to 100%?", lambda r: f'=V{r}+N(M{r})=100'),                                                                  # W
          ("Commissionable?", lambda r: f'=AND(D{r}="Closed Won",F{r}>={QS},F{r}<={QE},S{r},U{r},W{r})'),                          # X
          ("Credit (USD)", lambda r: f'=IF(X{r},G{r}*R{r}*T{r},0)'),                                                                # Y
          ("Fixed: owner credit", lambda r: f'=Y{r}*V{r}/100'),                                                                     # Z
          ("Fixed: split credit", lambda r: f'=Y{r}*N(M{r})/100'),                                                                  # AA
          ("SPIFF", lambda r: f'=IF(AND(X{r},K{r}="Yes",MONTH(F{r})=11,Y{r}>{SPMIN}),{SP},0)')]                                     # AB
header(ex, 1, [t for t, _ in T_COLS], start=15)
header(ex, 1, [t for t, _ in F_COLS], fill=RULE, start=17)
for r in range(2, N + 1):
    for i, (_, f) in enumerate(T_COLS + F_COLS):
        c = ex.cell(r, 15 + i, f(r)); c.font = F
        c.number_format = "yyyy-mm-dd" if i == 2 else ("0.0000" if i == 3 else "#,##0.00")
for i in range(1, 29):
    ex.column_dimensions[L(i)].width = 13
ex.column_dimensions["B"].width = 20; ex.freeze_panes = "B2"; ex.row_dimensions[1].height = 42

# ---------------- FX and Reps
fx = wb.create_sheet("FX"); header(fx, 1, ["week_of", "currency", "usd_per_unit"])
for r, row in enumerate(rows("fx_rates.csv"), 2):
    fx.cell(r, 1, d(row["week_of"])).number_format = "yyyy-mm-dd"; fx.cell(r, 2, row["currency"]); fx.cell(r, 3, float(row["usd_per_unit"]))
rp = wb.create_sheet("Reps")
header(rp, 1, ["rep", "start_date", "end_date", "Full months in quarter", "Quota", "Target incentive", "Employed on payment date?"], fill=RULE)
for r, row in enumerate(rows("reps.csv"), 2):
    rp.cell(r, 1, row["rep"]); rp.cell(r, 2, d(row["start_date"])).number_format = "yyyy-mm-dd"
    if row["end_date"]:
        rp.cell(r, 3, d(row["end_date"])).number_format = "yyyy-mm-dd"
    rp.cell(r, 4, f"=IF(B{r}<={QS},3,(B{r}<=DATE(2026,11,1))+(B{r}<=DATE(2026,12,1)))")
    rp.cell(r, 5, f"={QT}*D{r}/3").number_format = "#,##0"
    rp.cell(r, 6, f"={TI}*D{r}/3").number_format = "#,##0.00"
    rp.cell(r, 7, f'=OR(C{r}="",C{r}>={PAY})')
for c in "ABCDEFG":
    rp.column_dimensions[c].width = 16


# ---------------- the two payout tabs
def tiered(c, q, ti):
    return f"=MIN(3*{ti},{ti}/{q}*(MIN({c},{q})+1.5*MAX(0,MIN({c},1.25*{q})-{q})+2*MAX(0,{c}-1.25*{q})))"


def payout_tab(name, own, spl, fixed):
    ws = wb.create_sheet(name)
    ws["A1"] = f"Corvane Q4 FY2026 AE payouts, {name.lower()} spreadsheet"; ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws["A2"] = ("Credit = Closed Won in the quarter, by owner, split by split_pct. Tiered rate, capped at 3x target." if not fixed else
                "Same as Typical, plus every rule from the plan PDF, Tom's email and the Slack SPIFF (orange helper columns on Export and Reps).")
    ws["A2"].font = Font(name="Arial", size=10, italic=True)
    labels = ["Rep", "Quota", "Target incentive", "Credited bookings", "Attainment", "Commission"]
    labels += ["SPIFF", "Employed on Jan 15?", "Payout"] if fixed else ["Payout"]
    header(ws, 4, labels)
    for i, rep in enumerate(REPS):
        r, rr = 5 + i, 2 + i
        ws.cell(r, 1, rep)
        ws.cell(r, 2, f"=Reps!E{rr}" if fixed else f"={QT}")
        ws.cell(r, 3, f"=Reps!F{rr}" if fixed else f"={TI}")
        ws.cell(r, 4, f"=SUMIFS(Export!${own}$2:${own}${N},Export!$C$2:$C${N},A{r})+SUMIFS(Export!${spl}$2:${spl}${N},Export!$L$2:$L${N},A{r})")
        ws.cell(r, 5, f"=D{r}/B{r}")
        ws.cell(r, 6, tiered(f"D{r}", f"B{r}", f"C{r}"))
        if fixed:
            ws.cell(r, 7, f"=SUMIFS(Export!$AB$2:$AB${N},Export!$C$2:$C${N},A{r})")
            ws.cell(r, 8, f"=Reps!G{rr}")
            ws.cell(r, 9, f"=IF(H{r},F{r}+G{r},0)")
        else:
            ws.cell(r, 7, f"=F{r}")
        for c in range(1, len(labels) + 1):
            cell = ws.cell(r, c); cell.font = F
            if c > 1 and not (fixed and c == 8):
                cell.number_format = "0.0%" if c == 5 else "$#,##0.00"
    t, pc = 5 + len(REPS), len(labels)
    ws.cell(t, 1, "Total").font = B
    ws.cell(t, pc, f"=SUM({L(pc)}5:{L(pc)}{t-1})").font = B; ws.cell(t, pc).number_format = "$#,##0.00"
    ws.column_dimensions["A"].width = 16
    for c in range(2, pc + 1):
        ws.column_dimensions[L(c)].width = 16
    ws.row_dimensions[4].height = 30


payout_tab("Typical", "O", "P", False)
payout_tab("Fixed", "Z", "AA", True)
wb.move_sheet("Typical", -4); wb.move_sheet("Fixed", -4)
wb.save(OUT); print(OUT, [s.title for s in wb.worksheets])
